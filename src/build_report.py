"""Build the Russian report from completed experiments. No fabricated placeholders."""
from pathlib import Path
import json,csv,re,io,contextlib,datetime
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from update_literature import gost_reference
ROOT=Path(__file__).resolve().parents[1]
def j(path):return json.loads((ROOT/path).read_text())
def csvrows(path):
 with (ROOT/path).open(encoding='utf-8-sig') as f:return list(csv.DictReader(f))
def fmt(x,d=3):return '—' if x in (None,'') else f'{float(x):.{d}f}'.replace('.',',')
def main():
 done=j('results/analysis_complete.json');assert done['benchmark_responses']==2912
 sel=j('results/selection.json');cfg=j('configs/experiment.json');stats=j('results/data_analysis.json')['statistics'];matrix=j('literature/matrix.json')
 metrics=csvrows('results/model_metrics.csv');test=[r for r in metrics if r['split']=='test' and r['variant']=='main'];app=csvrows('results/application_metrics.csv');adapt=csvrows('results/adaptation/summary.csv');effects=csvrows('results/adaptation/paired_effects.csv');audit=j('results/adaptation/pseudo_label_audit.json');diffs=csvrows('results/paired_mode_comparisons.csv')
 doc=Document();sec=doc.sections[0];sec.page_height=Cm(29.7);sec.page_width=Cm(21);sec.top_margin=Cm(2);sec.bottom_margin=Cm(2);sec.left_margin=Cm(3);sec.right_margin=Cm(1.5)
 for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
  s=doc.styles[name];s.font.name='Times New Roman';s.font.color.rgb=RGBColor(0,0,0)
  fonts=s.element.get_or_add_rPr().rFonts
  if fonts is None:
   fonts=OxmlElement('w:rFonts');s.element.get_or_add_rPr().insert(0,fonts)
  for attr in ('ascii','hAnsi','eastAsia','cs'):fonts.set(qn('w:'+attr),'Times New Roman')
  # Built-in Word Title may carry a blue bottom border; academic report titles
  # use whitespace only.
  ppr=s.element.get_or_add_pPr()
  border=ppr.find(qn('w:pBdr'))
  if border is not None:ppr.remove(border)
 doc.styles['Normal'].font.size=Pt(14);doc.styles['Normal'].paragraph_format.line_spacing=1.5;doc.styles['Normal'].paragraph_format.first_line_indent=Cm(1.25);doc.styles['Normal'].paragraph_format.space_after=Pt(0);doc.styles['Normal'].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
 doc.styles['Title'].font.size=Pt(18);doc.styles['Title'].font.bold=True;doc.styles['Title'].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
 doc.styles['Heading 1'].font.size=Pt(14);doc.styles['Heading 1'].font.bold=True
 doc.styles['Heading 2'].font.size=Pt(14);doc.styles['Heading 2'].font.bold=True
 doc.styles['Heading 1'].paragraph_format.keep_with_next=True
 doc.styles['Heading 2'].paragraph_format.keep_with_next=True
 for name in ('Title','Heading 1','Heading 2'):
  doc.styles[name].paragraph_format.first_line_indent=Cm(0)
  doc.styles[name].paragraph_format.space_before=Pt(12)
  doc.styles[name].paragraph_format.space_after=Pt(6)
 footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
 fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
 doc.core_properties.title='Большие языковые модели в задачах адаптации визуального домена';doc.core_properties.subject='Лабораторная работа 1 на Office-Home';doc.core_properties.author='';doc.core_properties.keywords='Office-Home, MMD, CORAL, multimodal LLM, pseudo-labels';doc.core_properties.comments=''
 doc.core_properties.created=datetime.datetime.now(datetime.timezone.utc);doc.core_properties.modified=datetime.datetime.now(datetime.timezone.utc)
 def p(text):return doc.add_paragraph(text)
 def link(paragraph,label,url):
  from docx.opc.constants import RELATIONSHIP_TYPE
  relation=paragraph.part.relate_to(url,RELATIONSHIP_TYPE.HYPERLINK,is_external=True)
  el=OxmlElement('w:hyperlink');el.set(qn('r:id'),relation);run=OxmlElement('w:r');props=OxmlElement('w:rPr');color=OxmlElement('w:color');color.set(qn('w:val'),'1258A9');props.append(color);run.append(props);text=OxmlElement('w:t');text.text=label;run.append(text);el.append(run);paragraph._p.append(el)
 def h(text):doc.add_heading(text,2)
 table_count=[0]
 figure_count=[0]
 captions=[
  'Сравнительная матрица источников',
  'Состав полного набора Office-Home по доменам',
  'Разбиение исследовательского поднабора',
  'Семь мультимодальных моделей и число параметров',
  'Качество моделей на общем тесте Clipart',
  'Дополнительные показатели классификации на общем тесте',
  'Калибровка и время ответа моделей',
  'Чувствительность моделей к формулировке инструкции',
  'Результат дополнительной адаптации LoRA',
  'Качество выбранной модели на отдельных визуальных доменах',
  'Результаты сравнения способов выравнивания',
 ]
 def table(headers,rows,widths,keep_first_rows=False):
  table_count[0]+=1
  caption=doc.add_paragraph(f'Таблица {table_count[0]} — {captions[table_count[0]-1]}')
  caption.paragraph_format.first_line_indent=Cm(0);caption.paragraph_format.keep_with_next=True
  caption.paragraph_format.space_before=Pt(7);caption.paragraph_format.space_after=Pt(3)
  for run in caption.runs:run.font.size=Pt(12)
  t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
  factor=min(1.0,16.5/sum(widths));widths=[w*factor for w in widths]
  for c,w in zip(t.columns,widths):c.width=Cm(w)
  for c,s in zip(t.rows[0].cells,headers):c.text=str(s)
  repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
  for row in rows:
   for cell,value in zip(t.add_row().cells,row):cell.text=str(value)
  for ri,row in enumerate(t.rows):
   trpr=row._tr.get_or_add_trPr();nosplit=OxmlElement('w:cantSplit');trpr.append(nosplit)
   for ci,cell in enumerate(row.cells):
    cell.width=Cm(widths[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;pr=cell._tc.get_or_add_tcPr()
    borders=OxmlElement('w:tcBorders')
    for edge in ['top','left','bottom','right']:
     e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'6');e.set(qn('w:color'),'000000');borders.append(e)
    pr.append(borders);marg=OxmlElement('w:tcMar')
    for edge,value in [('top','45'),('bottom','45'),('left','90'),('right','90')]:
     e=OxmlElement('w:'+edge);e.set(qn('w:w'),value);e.set(qn('w:type'),'dxa');marg.append(e)
    pr.append(marg)
    for para in cell.paragraphs:
     para.paragraph_format.first_line_indent=Cm(0);para.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT
     para.paragraph_format.line_spacing=1.0;para.paragraph_format.space_after=Pt(1);para.paragraph_format.space_before=Pt(1)
     if keep_first_rows and ri in (0,1):para.paragraph_format.keep_with_next=True
     for run in para.runs:run.font.size=Pt(10.5);run.bold=ri==0
  spacer=doc.add_paragraph();spacer.paragraph_format.first_line_indent=Cm(0);spacer.paragraph_format.space_after=Pt(1)
 def figure(name,caption,width=16):
  figure_count[0]+=1
  para=doc.add_paragraph();para.paragraph_format.keep_with_next=True
  para.paragraph_format.first_line_indent=Cm(0)
  picture=para.add_run().add_picture(str(ROOT/'results/figures'/name),width=Cm(width))
  picture._inline.docPr.set('title',caption)
  picture._inline.docPr.set('descr',caption)
  para.alignment=WD_ALIGN_PARAGRAPH.CENTER
  cap=p(f'Рисунок {figure_count[0]} — {caption}');cap.paragraph_format.first_line_indent=Cm(0);cap.paragraph_format.space_after=Pt(10)
  for run in cap.runs:run.font.size=Pt(12);run.italic=True
 background=(ROOT/'report_background.md').read_text();parts=re.split(r'^## ',background,flags=re.M)
 # Title and scope, then dynamic abstract.
 for para in parts[0].strip().split('\n\n'):
  if para.startswith('# '):doc.add_paragraph(para[2:],'Title')
  else:p(para)
 win=next(r for r in test if r['model']==sel['model'] and r['mode']==sel['mode'])
 p(f'Краткий результат. Выполнено {done["benchmark_responses"]} ответов в основном сравнении семи моделей, {done["application_responses"]} ответов выбранной модели на отдельных данных и 18 запусков адаптации. По правилу source validation выбрана {sel["model"]}, режим {sel["mode"]}. На Clipart test её Macro-F1 составляет {fmt(win["macro_f1"])}; доля невалидных ответов — {fmt(100*float(win["invalid_rate"]),1)}%. Из 195 целевых изображений принято {audit["accepted"]} псевдометок. Полные значения и ограничения приведены далее.')
 if not sel['validity_gate_passed']:p('Порог допустимой доли невалидных ответов не пройден ни одним кандидатом. Выбранная модель применяется только для исследовательского сравнения; автоматическое использование в прототипе не рекомендовано.')
 for chunk in parts[1:]:
  title,body=chunk.split('\n',1);doc.add_heading(title.strip(),1)
  for para in body.strip().split('\n\n'):p(para)
  if title.startswith('Поиск'):
   para=p('Проверка рангов площадок по выпуску ICORE2026: ');link(para,'Официальный портал ICORE',j('literature/venue_ranks.json')['source'])
   h('Схема отбора источников')
   figure('literature_flow.png','Схема отбора зарегистрированных работ для аналитического обзора')
   short_roles=['данные','MMD','состязательное выравнивание','CORAL','условное выравнивание','визуально-языковые признаки','адаптация VLM','визуальные инструкции','компактная MLLM','семейство Qwen2-VL','семейство Qwen2.5-VL','адаптация LoRA','калибровка','псевдометки MLLM','язык инструкции']
   table(['№','Публикация','Метод или модель','Связь с работой'],[[str(i),r['authors']+' '+r['year'],r['architecture'].split(';')[0],short_roles[i-1]] for i,r in enumerate(matrix,1)],[1,3.5,8.2,3.8])
   p('Матрица охватывает 15 публикаций. Её полная версия с задачей, данными, протоколом, метриками, кодом, ограничениями и идеей репликации приложена как literature/matrix.csv.')
  if title.startswith('Подготовка'):
   table(['Домен','N','Min класс','Max класс','Дисбаланс'],[[s['domain'],s['n'],s['class_min'],s['class_max'],fmt(s['imbalance_ratio'],2)] for s in stats],[4.7,2,3,3,4])
   p('Дисбаланс — отношение максимального к минимальному числу объектов класса. Статистика относится к полному прочитанному корпусу, до отбора исследовательского поднабора.')
   table(['Часть','Домен','N','Доступ к меткам'],[['source_train','Product',520,'Обучение'],['validation','Product',130,'65 для LLM, 130 для классификатора'],['target_adapt','Clipart',195,'Только итоговый аудит'],['test','Clipart',130,'Только оценка'],['new_slice','Art',130,'Только оценка'],['external','Real_World',130,'Только оценка']],[3.5,3,1.5,8.7])
  if title.startswith('Единый'):
   h('Модели и версии')
   revisions=j('configs/model_revisions.json')
   table(['Ключ','Checkpoint','Параметры млрд'],[[m['key'],m['repo'],fmt(j(f'logs/{m["key"]}_benchmark_environment.json')['parameters']/1e9,3)] for m in cfg['models']],[3,11,2.7],keep_first_rows=True)
   p('Указано фактическое число параметров загруженного экземпляра, включая визуальную часть. Полные commit SHA находятся в configs/model_revisions.json; номинальный размер в названии может относиться только к языковой части.')
   p('По официальным карточкам SmolVLM, SmolVLM2 и Qwen2-VL-2B используют Apache-2.0; LLaVA Interleave Qwen 0.5B — Tongyi Qianwen Research License; Qwen2.5-VL-3B — Qwen Research License. Условия и ссылки зафиксированы в configs/model_cards_metadata.json. Отбор ограничен доступными для академического локального запуска checkpoint.')
   para=p('Первичная публикация, соответствующая использованному LLaVA-NeXT-Interleave checkpoint: F. Li et al., «LLaVA-NeXT-Interleave: Tackling Multi-image, Video, and 3D in Large Multimodal Models», 2024; ')
   link(para,'статья arXiv','https://arxiv.org/abs/2407.07895')
   para.add_run('. Сопоставление статьи и конкретной версии модели подтверждает ')
   link(para,'официальная карточка checkpoint','https://huggingface.co/llava-hf/llava-interleave-qwen-0.5b-hf')
   para.add_run('. Общая публикация LLaVA [8] остаётся архитектурной основой.')
   h('Качество на общем тесте')
   p('Значения в отчёте округлены до трёх знаков; малое ненулевое значение может отображаться как 0,000. Точные значения сохранены в CSV. Интервал [0,000; 0,000] на фиксированной выборке не доказывает нулевую точность на всей генеральной совокупности.')
   table(['Модель','Режим','Macro F1','95% ДИ','Accuracy','Invalid %'],[[r['model'],r['mode'],fmt(r['macro_f1']),fmt(r['macro_f1_ci_low'])+'–'+fmt(r['macro_f1_ci_high']),fmt(r['accuracy']),fmt(100*float(r['invalid_rate']),1)] for r in test],[3,2,2.1,4,2.1,2.1])
   table(['Модель','Режим','Weighted F1','Balanced acc.','MCC'],[[r['model'],r['mode'],fmt(r['weighted_f1']),fmt(r['balanced_accuracy']),fmt(r['mcc'])] for r in test],[3.5,2,3.5,4,3.5])
   h('Калибровка и вычислительные затраты')
   table(['Модель','Режим','Valid N','ECE','Brier','p50 с','p95 с'],[[r['model'],r['mode'],r['valid_n'],fmt(r['ece']),fmt(r['binary_brier']),fmt(r['latency_median_s'],2),fmt(r['latency_p95_s'],2)] for r in test],[3,2,2,2,2,2,2.4])
   p('Прочерк означает отсутствие валидных ответов, а не нулевую ошибку калибровки. Пиковый RSS и снимки памяти MPS приведены в results/memory_observations.csv; это разные показатели, их не суммируют. Неокруглённые метрики находятся в results/model_metrics.csv.')
   h('Разница few shot и zero shot')
   positives=sum(float(r['few_minus_zero_macro_f1'])>0 for r in diffs);reliable=sum(float(r['ci_low'])>0 for r in diffs)
   p(f'Положительная наблюдаемая разница получена для {positives} из 7 моделей; нижняя граница парного интервала выше нуля для {reliable} моделей. Это не подтверждает универсальную пользу few-shot. Парные разницы и интервалы для всех моделей приведены в results/paired_mode_comparisons.csv; множественные сравнения не корректировались.')
   h('Чувствительность инструкции')
   sensitivity=csvrows('results/prompt_sensitivity.csv')
   table(['Модель','Вариант','N','Смена метки %','Accuracy','Invalid %'],[[r['model'],r['variant'],r['n'],fmt(100*float(r['prediction_flip_rate']),1),fmt(r['accuracy']),fmt(100*float(r['invalid_rate']),1)] for r in sensitivity],[3,2,1,4.1,3,3.6])
   p('Все невалидные ответы кодируются −1. Поэтому нулевая смена метки при двух разных невалидных текстах не означает устойчивое следование инструкции. Для интерпретации вместе рассматриваются raw-ответы, accuracy и invalid rate.')
   chosen_sens=[r for r in sensitivity if r['model']==sel['model']]
   for r in chosen_sens:p(f'Для выбранной модели вариант {r["variant"]}: accuracy базовой инструкции на тех же 13 изображениях {fmt(r["base_accuracy_same_images"])}; нового варианта {fmt(r["accuracy"])}. Проверка использует zero-shot независимо от режима итогового выбора.')
   h('Выбор модели на validation')
   p(f'Зафиксирован выбор {sel["model"]}, {sel["mode"]}, S = {fmt(sel["selection_score"])}. Проверка порога валидности: '+('пройдена.' if sel['validity_gate_passed'] else 'не пройдена; применяется оговорённое исследовательское правило.'))
   p('Результаты всех кандидатов на Product validation, включая балл S, находятся в results/model_metrics.csv. Тестовый Clipart не использовался для выбора.')
   best_test=max(test,key=lambda r:float(r['macro_f1']))
   p(f'Максимальный наблюдаемый тестовый Macro-F1 получен у {best_test["model"]} в режиме {best_test["mode"]}: {fmt(best_test["macro_f1"])}. У выбранной по source validation модели и режима он равен {fmt(win["macro_f1"])}. Выбор после просмотра test не изменялся. Расхождение этих значений показывает ограничение переноса правила выбора с Product на Clipart.')
   h('Параметрически эффективная адаптация')
   peftfile=ROOT/'results/peft/protocol.json'
   if peftfile.exists() and j('results/peft/protocol.json').get('status')=='completed' and (ROOT/'results/peft/metrics.csv').exists():
    pp=j('results/peft/protocol.json');pm=csvrows('results/peft/metrics.csv')
    p(f'Дополнительный эксперимент LoRA выполнен для SmolVLM-500M: rank 8, alpha 16, dropout 0,05, q_proj/v_proj языковой части. Обучаемых параметров {pp["trainable_parameters"]:,} из {pp["total_parameters"]:,}. Две эпохи, 520 source_train, AdamW 10⁻⁴, накопление градиента 8, checkpoint по минимальной cross-entropy ответов на 65 source validation. Выбрана эпоха {pp["selected_epoch"]}. Длительность {fmt(pp["training_seconds"]/60,1)} мин; максимум RSS {fmt(pp["peak_process_rss_gb"],2)} GB.')
    table(['Режим','N test','Macro F1','Invalid %','ECE'],[[r['mode'],r['n'],fmt(r['macro_f1']),fmt(100*float(r['invalid_rate']),1),fmt(r['ece'])] for r in pm],[4,2,3,4,3.7])
    lora=next(r for r in pm if r['mode']=='lora')
    p(f'После LoRA правильно классифицировано {round(float(lora["accuracy"])*int(lora["n"]))} из {lora["n"]} тестовых изображений. Accuracy = {fmt(lora["accuracy"],4)}, Macro-F1 = {fmt(lora["macro_f1"],6)}. Улучшение формата ответа оценивается отдельно от содержательной точности разметки.')
    p('Cross-entropy обучения усредняется по всем токенам ответа. Большинство из них относится к постоянному JSON-шаблону и confidence=1,0; идентификатор класса занимает малую часть ответа. Поэтому уменьшение этой cross-entropy не равнозначно росту точности классификации. Возможный вклад такого дисбаланса токенов следует проверять отдельным экспериментом, например изменяя вес токенов label_id; причиной всех наблюдаемых ошибок он здесь не объявляется.')
    p('LoRA добавлена как отдельное разведочное расширение после начала основного сравнения. Она не участвует задним числом в выборе одной из семи базовых моделей. Обучающие ответы содержали confidence=1,0: это может усиливать чрезмерную уверенность и не является калибровочным обучением. Истинные target-метки не использовались.')
   else:
    raise RuntimeError('PEFT extension has not been resolved; do not finalize report')
   h('Ошибки и предметные срезы')
   p('На test встречались ответы в блоке Markdown, JSON не в виде объекта, ошибки схемы и синтаксиса. Категории взаимно исключаются; их частоты приведены в results/schema_errors.csv. Ответы не исправлялись для метрик или выбора модели.')
   p('Анализ ошибок по классам сохранён в results/per_class.csv, по длине входных и выходных токенов, языку инструкции, площади исходного изображения и домену — в results/error_strata.csv. В Office-Home нет финансовых сообщений, поэтому соответствующий пункт задания интерпретируется как тип визуального домена. Длины и размеры рассматриваются описательно: они не рандомизированы, группы могут различаться составом объектов. В языковых сравнениях одни изображения повторяются, их нельзя считать независимыми.')
   class_rows=[r for r in csvrows('results/per_class.csv') if r['model']==sel['model'] and r['mode']==sel['mode']]
   p(f'Для выбранной модели и режима хотя бы одно правильное предсказание наблюдается в {sum(int(r["correct"])>0 for r in class_rows)} из 65 классов; суммарно верны {sum(int(r["correct"]) for r in class_rows)} из 130 изображений test. Это показывает, почему одной общей accuracy недостаточно для анализа распределения ошибок по классам.')
   p('Срезы по площади изображения и длине ответа доступны в results/error_strata.csv. Их различия описательные: размеры связаны с составом объектов, поэтому причину ошибки по ним установить нельзя. Влияние языка оценено на 13 повторно использованных изображениях.')
   es=j('results/selected_error_summary.json')
   p(f'У выбранной модели {es["valid_errors"]} валидных ошибочных ответов; {es["confident_wrong_answers"]} из них имеют confidence не ниже {fmt(es["confidence_threshold"],1)}. Наличие таких ответов показывает, что установленный порог уверенности не гарантирует точности псевдометки. По этим тестовым ошибкам порог не перенастраивался.')
   p('Наиболее частые пары ошибочных классов указаны в results/selected_error_summary.json. При двух изображениях каждого истинного класса эти частоты неустойчивы. Десять примеров ниже выбраны из разных классов и не предназначены для оценки распространённости ошибок.')
   figure('ten_errors.png','Десять реальных ошибок с идентификаторами исходных изображений')
   p('Идентификаторы, настоящие и предсказанные классы видны на рисунке; исходные ответы и типы всех десяти ошибок содержатся в results/ten_errors.json. Полная confusion matrix с колонкой INVALID приложена в CSV и PNG.')
  if title.startswith('Связь'):
   h('Применение выбранной модели к отдельным данным')
   table(['Срез','N','Macro F1','Accuracy','Invalid %'],[[r['domain'],r['n'],fmt(r['macro_f1']),fmt(r['accuracy']),fmt(100*float(r['invalid_rate']),1)] for r in app],[4.7,2,3,3,4])
   p(f'Принято {audit["accepted"]} из {audit["n_target_adapt"]} псевдометок. Точность принятых меток при отдельном аудите: {fmt(audit["accepted_accuracy_for_evaluation_only"])}. Эта проверка не использовалась для изменения порога 0,7 или подбора методов.')
   top=audit['most_frequent_predictions'][0] if audit['most_frequent_predictions'] else None
   if top:
    p(f'Правильных принятых меток: {audit["accepted_correct_for_evaluation_only"]}. В принятых предсказаниях представлены {audit["predicted_classes_present"]} из 65 классов; наиболее частый класс — {top["class"]}, {top["count"]} ответов из {audit["accepted"]}. Преобладание одного класса показывает смещение распределения псевдометок относительно сбалансированной целевой выборки.')
    if top['label_id'] in [0,22,43]:p('Этот класс представлен среди трёх few-shot демонстраций. Возможное влияние примеров остаётся гипотезой: отдельная проверка состава демонстраций здесь не выполнена.')
   h('Результаты адаптации')
   domain_name={'test':'Clipart','new_slice':'Art'}
   method_name={'none':'Без выравн.','mmd':'MMD','coral':'CORAL'}
   table(['Домен оценки','Метод','Псевдометки','Macro F1 ± SD','Accuracy ± SD'],[[domain_name.get(r['split'],r['split']),method_name.get(r['alignment'],r['alignment']),('да' if r['pseudo_labels']=='True' else 'нет'),fmt(r['macro_f1_mean'])+' ± '+fmt(r['macro_f1_sd']),fmt(r['accuracy_mean'])+' ± '+fmt(r['accuracy_sd'])] for r in adapt],[3.2,2.3,3,4.1,4.1])
   h('Вклад выравнивания относительно общего baseline')
   alignment_effects=csvrows('results/adaptation/alignment_effects.csv')
   p('Вклад выравнивания относительно общего baseline рассчитан попарно при одинаковом seed и наличии псевдометок; полная таблица находится в results/adaptation/alignment_effects.csv. Следующий показатель сравнивает добавление псевдометок при неизменном методе.')
   for e in effects:
    if e['split']=='test':p(f'{method_name.get(e["alignment"],e["alignment"])}: средняя парная разница Macro-F1 от добавления псевдометок {fmt(e["pseudo_effect_mean"])}; SD разницы между seed {fmt(e["pseudo_effect_sd"])}.')
   p('Знак разницы отражает наблюдаемый эффект на этом поднаборе. Три seed и один target-домен не дают оснований утверждать универсальное превосходство метода или статистическую значимость только по знаку среднего.')
 doc.add_heading('Вывод по исследовательскому вопросу',1)
 p(f'В заданном протоколе многокритериальное правило выбрало {sel["model"]} в режиме {sel["mode"]}. Тестовый Macro-F1 равен {fmt(win["macro_f1"])}, интервал {fmt(win["macro_f1_ci_low"])}–{fmt(win["macro_f1_ci_high"])}. Это результат конкретной версии модели, инструкции и поднабора, а не общий рейтинг мультимодальных LLM.')
 p('Польза few-shot оказалась зависимой от модели: положительное изменение наблюдалось у '+str(sum(float(r['few_minus_zero_macro_f1'])>0 for r in diffs))+' из семи checkpoint. Корректность JSON и уверенность необходимо проверять отдельно от содержательного распознавания.')
 p(f'Гипотеза о полезности формально корректных уверенных псевдометок не подтверждается самим прохождением порога: среди принятых меток target_adapt правильны {audit["accepted_correct_for_evaluation_only"]} из {audit["accepted"]}. Для автоматической разметки нового домена требуется независимая проверка содержательной точности; высокий confidence в этом опыте её не заменяет.')
 baseline_methods=[r for r in adapt if r['split']=='test' and r['pseudo_labels']=='False']
 best_alignment=max(baseline_methods,key=lambda r:float(r['macro_f1_mean']))
 p(f'В сравнении без внешних псевдометок наибольший наблюдаемый средний Macro-F1 получен у {best_alignment["alignment"]}: {fmt(best_alignment["macro_f1_mean"])} ± {fmt(best_alignment["macro_f1_sd"])} SD. Это описание данного пилотного результата; статистически доказанное или универсальное превосходство метода не заявляется.')
 p('Для сравнения методов выравнивания подтверждены следующие измеренные изменения после добавления псевдометок: '+ '; '.join(method_name.get(e['alignment'],e['alignment'])+' '+fmt(e['pseudo_effect_mean'])+' Macro-F1' for e in effects if e['split']=='test')+'. Гипотеза о безусловном улучшении от внешней разметки из этих результатов не следует. Тема ВКР сохраняется: дальнейший анализ должен оценивать методы выравнивания в одинаковых условиях и отделять их эффект от качества внешних меток.')
 doc.add_heading('Воспроизводимые артефакты',1)
 p('В комплект входят исходный код, notebook анализа, README, матрица литературы и BibTeX, паспорт и аудит данных, manifest, конфигурации запросов, версии checkpoint, сырые ответы, таблицы метрик, графики и десять ошибок. Порядок повторения приведён в README.md. Скачанные веса после расчётов удаляются; их ревизии сохраняются для повторной загрузки.')
 remote=ROOT/'configs/repository_url.txt'
 if remote.exists():
  url=remote.read_text().strip()
  para=p('Репозиторий студента: ')
  link(para,url,url)
 else:p('Код предоставлен в локальном репозитории lab1_officehome. Удалённая ссылка отсутствует; для формального выполнения пункта 9 требований к отчёту требуется публикация этого репозитория и добавление реальной ссылки.')
 doc.add_heading('Список литературы',1)
 p('Нумерация записей соответствует ссылкам в тексте и строкам сравнительной матрицы. Описания содержат подтверждённые выходные данные и адреса первоисточников; дата обращения к электронным версиям — 9 октября 2026 года.')
 citations=j('literature/citation_metadata.json')
 for i,(r,m) in enumerate(zip(matrix,citations),1):
  assert r['id']==m['id'] and r['url']==m['url']
  record=gost_reference(m,r['title'])
  prefix,rest=record.split('URL: ',1)
  url,suffix=rest.split(' (дата обращения',1)
  para=p(f'[{i}] '+prefix+'URL: ')
  para.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT
  para.paragraph_format.first_line_indent=Cm(0)
  para.paragraph_format.left_indent=Cm(0.8)
  para.paragraph_format.first_line_indent=Cm(-0.8)
  link(para,url,url)
  para.add_run(' (дата обращения'+suffix)
 assert table_count[0]==len(captions),(table_count[0],len(captions))
 target=ROOT/'Лабораторная_1_Office_Home.docx';doc.save(target);print(target)
if __name__=='__main__':main()
